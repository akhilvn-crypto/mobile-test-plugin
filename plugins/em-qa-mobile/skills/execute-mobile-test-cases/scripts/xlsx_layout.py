"""Shared XLSX layout rules: wrapped text, clean alignment, rows tall enough for their content.

Every script that writes or updates an XLSX in this skill pair runs the matching tidy_* function last,
so no cell overflows into its neighbour and no wrapped text is clipped. A copy of this file lives in
both skills (build-mobile-test-cases/scripts and execute-mobile-test-cases/scripts); keep the two identical.

openpyxl cannot auto-fit rows and Excel does not re-fit them on open, so heights are computed here
from the text, glyph widths of the cell's font (Calibri / Arial metrics), the column width (merged
ranges use their summed width) and the font size. Calibrated against Excel's own AutoFit: no row comes
out shorter than Excel would make it.

Sheet rules (all idempotent, safe to run on a finished workbook):
  tidy_testcase_sheet(ws)  "Test Case ID" sheet: readable widths by header name, header wrapped and
                           centred, data wrapped/top/left (status, people, date, defects centred)
  tidy_report_sheet(ws)    "Test Report" sheet: widths, label merges so nothing spills, header rows
                           centred, data wrapped/top (#, Testcase #, Status centred)
  tidy_info_sheet(ws)      Document Control / Version History: wrap every filled cell, widen values
  tidy_workbook(wb)        apply the right rule to every sheet

CLI, to tidy an existing workbook in place (a backup goes to <folder>/.history/ first):
  python xlsx_layout.py "<file.xlsx>" [more files...] [--no-backup]
"""
import math
import os
import re
import shutil
import sys

from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string, get_column_letter

LINE_PT = 1.25      # row points per line, per font point (calibrated against Excel AutoFit)
PAD_PT = 6          # breathing room above/below the text (print renders a little tighter)
MAX_ROW_PT = 409    # Excel's limit
SAFETY = 0.92       # use only this share of a line's width, so estimates err towards taller rows
CELL_PAD = 0.9      # column-width units lost to the cell's inner padding

# Advance widths (font units / 2048) of common glyphs, used to measure text in "width of '0'" units,
# which is the unit Excel column widths are expressed in.
_GLYPHS = "0abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ .,-\"()':/;!?&%"
_METRICS = {
    "calibri": dict(zip(_GLYPHS, [
        1038, 981, 1076, 866, 1076, 1019, 625, 964, 1076, 470, 490, 931, 470, 1636, 1076, 1080, 1076,
        1076, 714, 801, 686, 1076, 925, 1464, 887, 927, 809, 1185, 1114, 1092, 1260, 1000, 941, 1292,
        1276, 516, 653, 1064, 861, 1751, 1322, 1356, 1058, 1378, 1112, 941, 998, 1314, 1162, 1822,
        1063, 998, 959, 463, 517, 511, 631, 820, 621, 621, 452, 548, 792, 548, 546, 941, 1400, 1465])),
    "arial": dict(zip(_GLYPHS, [
        1139, 1139, 1139, 1024, 1139, 1139, 569, 1139, 1139, 455, 455, 1024, 455, 1706, 1139, 1139, 1139,
        1139, 682, 1024, 569, 1139, 1024, 1479, 1024, 1024, 1024, 1366, 1366, 1479, 1479, 1366, 1251,
        1593, 1479, 569, 1024, 1366, 1139, 1706, 1479, 1593, 1366, 1593, 1479, 1366, 1251, 1479, 1366,
        1933, 1366, 1366, 1251, 569, 569, 569, 682, 727, 682, 682, 391, 569, 569, 569, 569, 1139, 1366,
        1821])),
}


def _char_widths(font_name):
    name = (font_name or "").lower()
    is_arial = any(k in name for k in ("arial", "helvetica", "liberation sans"))
    table = _METRICS["arial" if is_arial else "calibri"]
    zero = table["0"]
    upper_default = max(table[c] for c in "ABCDEFGHIJKLMNOPQRSTUVWXYZ") / zero
    return {c: w / zero for c, w in table.items()}, upper_default, 1.0


def _text_width(text, widths, upper_default, other_default):
    return sum(widths.get(ch, upper_default if ch.isupper() else other_default) for ch in text)


def _merged_lookup(ws):
    """top-left coordinate -> (total width, spans_rows); other cells in a merge -> None."""
    info = {}
    for m in ws.merged_cells.ranges:
        w = sum(col_width(ws, c) for c in range(m.min_col, m.max_col + 1))
        info[(m.min_row, m.min_col)] = (w, m.max_row > m.min_row)
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col):
                    info[(r, c)] = None
    return info


def col_width(ws, col):
    letter = get_column_letter(col) if isinstance(col, int) else col
    dim = ws.column_dimensions.get(letter)  # .get: do not create an empty dimension
    if dim is not None and dim.width:
        return dim.width
    # a dimension can cover a range (min..max); look for one that contains this column
    idx = col if isinstance(col, int) else column_index_from_string(letter)
    for d in ws.column_dimensions.values():
        if d.min and d.max and d.min <= idx <= d.max and d.width:
            return d.width
    return ws.sheet_format.defaultColWidth or 8.43


def set_widths(ws, widths, only_grow=False):
    for letter, w in widths.items():
        cur = col_width(ws, letter)
        if only_grow and cur and cur >= w:
            continue
        ws.column_dimensions[letter].width = w


def _with_alignment(cell, **kw):
    a = cell.alignment
    base = dict(horizontal=a.horizontal, vertical=a.vertical, wrap_text=a.wrap_text, indent=a.indent,
                text_rotation=a.text_rotation, shrink_to_fit=False)
    base.update(kw)
    cell.alignment = Alignment(**base)


def style_header(ws, row, cols):
    for c in cols:
        _with_alignment(ws.cell(row, c), horizontal="center", vertical="center", wrap_text=True)


def style_body(ws, rows, cols, center=(), vertical="top"):
    center = set(center)
    for r in rows:
        for c in cols:
            _with_alignment(ws.cell(r, c), horizontal="center" if c in center else "left",
                            vertical=vertical, wrap_text=True)


def wrap_values(ws, cells, vertical="top"):
    for cell in cells:
        _with_alignment(cell, vertical=vertical, wrap_text=True)


def _lines(text, capacity, measure):
    """Greedy word wrap, the way Excel does it; returns the number of lines."""
    space = measure(" ")
    total = 0
    for para in str(text).split("\n"):
        if not para.strip():
            total += 1
            continue
        n, cur = 1, 0.0
        for word in para.split(" "):
            wl = measure(word)
            if cur and cur + space + wl <= capacity:
                cur += space + wl
                continue
            if cur:
                n += 1
            if wl > capacity:  # a long unbroken word (a URL) is split across lines
                n += math.ceil(wl / capacity) - 1
                cur = wl % capacity or capacity
            else:
                cur = wl
        total += n
    return total


def _default_font_size(ws):
    """A cell font without a size inherits the workbook's default (Arial 10 in Sheets-made templates)."""
    try:
        return ws.parent._fonts[0].sz or 11
    except (AttributeError, IndexError):
        return 11


def needed_height(ws, r, merged=None):
    merged = _merged_lookup(ws) if merged is None else merged
    default_size = _default_font_size(ws)
    need = 0.0
    for cell in ws[r]:
        if cell.value is None or (isinstance(cell.value, str) and cell.value.startswith("=")):
            continue
        key = (cell.row, cell.column)
        if key in merged:
            if merged[key] is None or merged[key][1]:
                continue  # hidden part of a merge, or a merge spanning rows
            width = merged[key][0]
        else:
            width = col_width(ws, cell.column)
        size = cell.font.sz if cell.font and cell.font.sz else default_size
        bold = bool(cell.font and cell.font.b)
        widths, up, other = _char_widths(cell.font.name if cell.font else None)
        scale = (size / default_size) * (1.07 if bold else 1.0)

        def measure(t, widths=widths, up=up, other=other, scale=scale):
            return _text_width(t, widths, up, other) * scale

        capacity = max(1.0, (width - CELL_PAD) * SAFETY)
        if cell.alignment is not None and cell.alignment.wrap_text:
            lines = _lines(cell.value, capacity, measure)
        else:
            lines = str(cell.value).count("\n") + 1
        need = max(need, lines * size * LINE_PT + PAD_PT)
    return min(need, MAX_ROW_PT)


def fit_rows(ws, rows=None, min_height=None, shrink=False):
    """Size each row to its wrapped content. Grows only, unless shrink=True (rows this pipeline owns)."""
    merged = _merged_lookup(ws)
    rows = rows if rows is not None else range(1, ws.max_row + 1)
    for r in rows:
        need = needed_height(ws, r, merged)
        if not need:
            continue
        cur = ws.row_dimensions[r].height
        floor = max(15 if shrink else (cur or 15), min_height or 0)
        target = round(max(need, floor), 2)
        if cur is None or abs(target - cur) > 0.5:
            ws.row_dimensions[r].height = target


# ------------------------------------------------------------------ sheet rules

def _norm(v):
    v = re.sub(r"\(.*?\)", " ", str(v or "").replace("\n", " ")).lower()
    return re.sub(r"\s+", " ", v).strip()


# minimum widths per Testcases header (a template width wins when it is wider)
TESTCASE_WIDTHS = {"test case id": 18, "test case title": 30, "test description": 38, "preconditions": 30,
                   "test steps": 48, "test data": 30, "expected result": 44, "actual result": 34,
                   "execution status": 13, "comments": 30, "created by": 12, "reviewed by": 12,
                   "executed by": 12, "executed on": 22, "execution defects": 14}
TESTCASE_CENTER = {"execution status", "created by", "reviewed by", "executed by", "executed on",
                   "execution defects"}
REPORT_WIDTHS = {"A": 13, "B": 14, "C": 30, "D": 46, "E": 20, "F": 14, "G": 40}
REPORT_CENTER = (1, 5, 6)  # #, Testcase #, Status
REPORT_LABEL_MERGES = ("B2:D2", "A3:C3")  # TestSuite name; "Executed by:" (row 4 is merged A:C already)
INFO_WIDTHS = {"Document Control": {"D": 26, "E": 60},
               "Version History": {"E": 12, "F": 14, "G": 18, "H": 18, "I": 45}}


def find_header(ws, text, max_row=10):
    t = _norm(text)
    for r in range(1, min(ws.max_row, max_row) + 1):
        for c in range(1, min(ws.max_column, 40) + 1):
            if _norm(ws.cell(r, c).value) == t:
                return r, c
    return None, None


def _band_rows(ws):
    """Rows holding a single-row merge from column A (flow bands): template-styled, only height-fitted."""
    return {m.min_row for m in ws.merged_cells.ranges if m.max_row == m.min_row and m.min_col == 1}


def tidy_testcase_sheet(ws):
    hrow, _ = find_header(ws, "Test Case ID")
    if hrow is None:
        return False
    cols, center = {}, []
    for c in range(1, ws.max_column + 1):
        k = _norm(ws.cell(hrow, c).value)
        if not k:
            continue
        cols[c] = k
        if k in TESTCASE_CENTER:
            center.append(c)
    set_widths(ws, {get_column_letter(c): TESTCASE_WIDTHS[k] for c, k in cols.items() if k in TESTCASE_WIDTHS},
               only_grow=True)
    style_header(ws, hrow, cols)
    bands = _band_rows(ws)
    last = max((r for r in range(hrow + 1, ws.max_row + 1)
                if any(ws.cell(r, c).value is not None for c in cols)), default=hrow)
    body = [r for r in range(hrow + 1, last + 1) if r not in bands]
    style_body(ws, body, cols, center=center)
    fit_rows(ws, [hrow] + sorted(bands))
    fit_rows(ws, body, shrink=True)
    return True


def _merge_if_free(ws, rng):
    from openpyxl.worksheet.cell_range import CellRange
    want = CellRange(rng)
    if any(not want.isdisjoint(m) for m in ws.merged_cells.ranges):
        return
    rest = [ws.cell(r, c) for r in range(want.min_row, want.max_row + 1)
            for c in range(want.min_col, want.max_col + 1)][1:]
    if all(c.value is None for c in rest):
        ws.merge_cells(rng)


def tidy_report_sheet(ws, ncols=7):
    set_widths(ws, REPORT_WIDTHS, only_grow=True)
    for rng in REPORT_LABEL_MERGES:
        _merge_if_free(ws, rng)
    headers = [r for r in range(1, ws.max_row + 1) if _norm(ws.cell(r, 1).value) == "#"]
    first_band = (headers[0] - 1) if headers else ws.max_row + 1
    # summary block above the first band: wrap, keep the template's vertical centring
    wrap_values(ws, [ws.cell(r, c) for r in range(1, first_band)
                     for c in range(1, ncols + 1) if ws.cell(r, c).value is not None], vertical="center")
    body = []
    for h in headers:
        style_header(ws, h, range(1, ncols + 1))
        r = h + 1
        while r <= ws.max_row and any(ws.cell(r, c).value is not None for c in range(1, ncols + 1)):
            body.append(r)
            r += 1
    style_body(ws, body, range(1, ncols + 1), center=REPORT_CENTER)
    wrap_values(ws, [ws.cell(h - 1, ncols) for h in headers], vertical="center")  # "Execution Date : ..."
    owned = set(body)
    fit_rows(ws, [r for r in range(1, ws.max_row + 1) if r not in owned])
    fit_rows(ws, body, shrink=True)
    return True


def tidy_info_sheet(ws, min_widths=None):
    set_widths(ws, min_widths or INFO_WIDTHS.get(ws.title, {}), only_grow=True)
    wrap_values(ws, [c for row in ws.iter_rows() for c in row if c.value is not None])
    fit_rows(ws)


def tidy_workbook(wb):
    done = []
    for ws in wb.worksheets:
        if ws.title in INFO_WIDTHS:
            tidy_info_sheet(ws)
            done.append(ws.title)
        elif ws.title.strip().lower() in ("test report", "report template"):
            tidy_report_sheet(ws)
            done.append(ws.title)
        elif tidy_testcase_sheet(ws):
            done.append(ws.title)
    return done


def main(argv):
    from datetime import datetime, timedelta, timezone

    from openpyxl import load_workbook
    files = [a for a in argv if not a.startswith("--")]
    if not files:
        sys.exit(__doc__)
    for path in files:
        if not os.path.isfile(path):
            print(f"not found: {path}")
            continue
        if "--no-backup" not in argv:
            hist = os.path.join(os.path.dirname(os.path.abspath(path)), ".history")
            os.makedirs(hist, exist_ok=True)
            ts = datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%Y-%m-%d_%H-%M-%S-IST")
            stem, ext = os.path.splitext(os.path.basename(path))
            shutil.copy2(path, os.path.join(hist, f"{stem}__{ts}{ext}"))
        wb = load_workbook(path)
        sheets = tidy_workbook(wb)
        tmp = path + ".tidy.xlsx"
        wb.save(tmp)
        os.replace(tmp, path)
        print(f"tidied {path}: {', '.join(sheets) or 'no known sheets'}")


if __name__ == "__main__":
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    main(sys.argv[1:])
