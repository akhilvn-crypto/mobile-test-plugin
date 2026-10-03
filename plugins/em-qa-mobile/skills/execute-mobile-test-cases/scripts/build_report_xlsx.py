#!/usr/bin/env python3
"""Build the XLSX Test Report from the user's template (execute-mobile-test-cases Phase 7, only with
--test-report xlsx). Platform, device, OS and app build go into the Document Control Description and the Version
History comment; nothing else in the structure changes.

  python build_report_xlsx.py --data "Test Reports/.data/<report>.json" [--out "<report>.xlsx"] [--no-libreoffice]

The data file is written by build_report_md.py, so the MD and XLSX reports always show the same numbers.
--out defaults to the MD report's name with .xlsx.

How the template is followed (assets/Test_Report_Template.xlsx):
- The workbook is copied; the 4th sheet "Report templates" is deleted; "Report Template" is renamed "Test Report".
- Document Control: same rows; sample values cleared; only Title, Description ("Test report for <plan>") and
  Privacy Classification ("Confidential") are filled.
- Version History: same columns; sample rows cleared; one row (version, created on, comment).
- Report sheet: title band unchanged; TestSuite = plan name; Execution Date = run date; "Executed by" and
  "Overall testing by lead" empty; summary formulas kept (range end extended if needed); one green
  TestSuite band per flow with the template's header row, a blank row between flows; the status dropdown and
  colour rules are rebuilt explicitly over every status cell.
- Layout: xlsx_layout.tidy_workbook() runs last on every sheet - readable column widths, label cells
  merged so they never spill, headers wrapped and centred, data wrapped/top-aligned (#, Testcase # and
  Status centred), every row tall enough for its text so nothing overflows or is clipped.
- Comments: a bug filed in Jira shows as "PROJ-45 (BUG-003) - <title>" with the issue as the cell's hyperlink
  (only when the row names one Jira issue; a cell can hold one link).
- Formulas: recalculated with LibreOffice when it is installed, otherwise the workbook recalculates on load.
"""
import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
from copy import copy

from openpyxl.styles.cell_style import StyleArray

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import dump, find_root, load_json, rel, utf8_stdio  # noqa: E402
import xlsx_layout  # noqa: E402

SKILL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(SKILL, "assets", "Test_Report_Template.xlsx")
STATUS_LIST = ["PASS", "FAIL", "IN PROGRESS", "UNEXECUTED", "BLOCKED"]
STATUS_FILL = {"PASS": "FFB7E1CD", "FAIL": "FFCC4125", "IN PROGRESS": "FF4A86E8", "BLOCKED": "FFFFFF00",
               "UNEXECUTED": "FFFFFFFF"}
REPORT_SHEET = "Test Report"
FIRST_BAND_ROW = 13  # template: band 13, header 14, first data row 15
COLS = 7  # A..G


def find_label(ws, text):
    t = text.strip().lower()
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip().lower().rstrip(":").strip() == t:
                return c
    return None


def fill_document_control(ws, data):
    labels = ["Title", "Document ID", "Description", "Author", "Review Date", "Reviewed by", "Approved date",
              "Approved By", "Privacy Classification"]
    desc = f"Test report for {data['plan_name']}" + (f", {data['mobile_line']}" if data.get("mobile_line") else "")
    values = {"Title": data["title"], "Description": desc,
              "Privacy Classification": "Confidential"}
    for lab in labels:
        c = find_label(ws, lab)
        if c is None:
            continue
        ws.cell(c.row, c.column + 1).value = values.get(lab)


def fill_version_history(ws, data):
    head = find_label(ws, "Version No")
    if head is None:
        raise SystemExit("Version History: 'Version No' header not found in the template")
    r0, c0 = head.row, head.column
    first_style = [copy(ws.cell(r0 + 1, c0 + i)._style) for i in range(5)]
    for r in range(r0 + 1, ws.max_row + 1):
        for i in range(5):
            cell = ws.cell(r, c0 + i)
            cell.value = None
            if r > r0 + 1:
                cell._style = StyleArray()  # plain style
    comment = data["version_comment"] + (f" ({data['mobile_line']})" if data.get("mobile_line") else "")
    vals = [data["version"], data["created_on"], None, None, comment]
    for i, v in enumerate(vals):
        cell = ws.cell(r0 + 1, c0 + i)
        cell._style = copy(first_style[i])
        cell.number_format = "@" if i == 1 else cell.number_format
        cell.value = v


def capture_row(ws, r):
    return {"height": ws.row_dimensions[r].height,
            "cells": [(ws.cell(r, c).value, copy(ws.cell(r, c)._style)) for c in range(1, COLS + 1)]}


def apply_style(ws, r, captured, values=None):
    if captured["height"]:
        ws.row_dimensions[r].height = captured["height"]
    for c in range(1, COLS + 1):
        cell = ws.cell(r, c)
        cell._style = copy(captured["cells"][c - 1][1])
        if values is not None:
            cell.value = values[c - 1]


def fill_report_sheet(ws, data):
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList

    band = capture_row(ws, FIRST_BAND_ROW)
    header = capture_row(ws, FIRST_BAND_ROW + 1)
    datarow = capture_row(ws, FIRST_BAND_ROW + 2)
    plain = StyleArray()
    old_rules = []
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            old_rules.append(rule)

    # header area
    ws["B2"] = data["plan_name"]
    exec_label = find_label(ws, "Execution Date")
    if exec_label is not None and exec_label.row == 2:
        ws.cell(2, exec_label.column + 1).value = data["execution_date"]
    for r in (3, 4):
        for c in range(2, COLS + 1):
            cell = ws.cell(r, c)
            if type(cell).__name__ == "MergedCell":
                continue
            if not isinstance(cell.value, str) or not cell.value.strip().endswith(":"):
                cell.value = None

    # remove the sample bands
    for m in [m for m in ws.merged_cells.ranges if m.min_row >= FIRST_BAND_ROW]:
        ws.unmerge_cells(str(m))
    last_template_row = ws.max_row
    for r in range(FIRST_BAND_ROW, last_template_row + 1):
        ws.row_dimensions[r].height = None
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell._style = copy(plain)

    # one band per flow
    r = FIRST_BAND_ROW
    status_ranges = []
    for f in data["flows"]:
        if not f["rows"]:
            continue
        apply_style(ws, r, band, [band["cells"][0][0] or "TestSuite : ", f["name"], None, None, None, None,
                                  f"Execution Date : {data['execution_date']}"])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        apply_style(ws, r + 1, header, [v for v, _ in header["cells"]])
        first = r + 2
        for i, row in enumerate(f["rows"]):
            rr = first + i
            apply_style(ws, rr, datarow, [row["n"], row["story_ref"] or None, row["story_desc"], row["scenario"],
                                          row["id"], row["status"], row["comment"] or None])
            urls = sorted(set(row.get("links") or []))
            if len(urls) == 1:
                ws.cell(rr, 7).hyperlink = urls[0]
            fill = STATUS_FILL.get(row["status"])
            if fill:
                ws.cell(rr, 6).fill = PatternFill(fill_type="solid", fgColor=fill, bgColor=fill)
        last = first + len(f["rows"]) - 1
        status_ranges.append(f"F{first}:F{last}")
        r = last + 2  # one blank row between flows
    last_row = max(r, 1001)

    # summary formulas: keep them, extend their range when the report is longer than the template's
    for rr in range(5, 12):
        cell = ws.cell(rr, 6)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.value = re.sub(r"\$F\$13:\$F\d+", f"$F${FIRST_BAND_ROW}:$F{last_row}", cell.value)

    # dropdown and colour rules over every status cell
    ws.data_validations = DataValidationList()
    ws.conditional_formatting = ConditionalFormattingList()
    if status_ranges:
        sqref = " ".join(status_ranges)
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_LIST) + '"', allow_blank=True)
        for rng in status_ranges:
            dv.add(rng)
        ws.add_data_validation(dv)
        top_left = status_ranges[0].split(":")[0]
        for rule in old_rules:
            rule = copy(rule)
            if rule.formula:
                rule.formula = [re.sub(r"\bF\d+\b", top_left, f) for f in rule.formula]
            ws.conditional_formatting.add(sqref, rule)
    return status_ranges



def find_soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(p):
            return p
    return None


def recalc_with_libreoffice(path):
    soffice = find_soffice()
    if not soffice:
        return False
    out_dir = tempfile.mkdtemp(prefix="report-recalc-")
    try:
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                        "--outdir", out_dir, path], check=True, capture_output=True, timeout=180)
        produced = os.path.join(out_dir, os.path.basename(path))
        if os.path.exists(produced):
            shutil.copyfile(produced, path)
            return True
    except (subprocess.SubprocessError, OSError):
        return False
    finally:
        shutil.rmtree(out_dir, ignore_errors=True)
    return False


def main():
    utf8_stdio()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", required=True)
    ap.add_argument("--out")
    ap.add_argument("--no-libreoffice", action="store_true")
    a = ap.parse_args()
    root = find_root()
    import openpyxl
    data = load_json(a.data)
    if not data:
        dump({"ok": False, "error": f"Report data not found: {a.data}"}, code=2)
    stem = os.path.splitext(os.path.basename(a.data))[0]
    out = a.out or os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(a.data))), stem + ".xlsx")
    if os.path.exists(out):
        dump({"ok": False, "error": f"{out} already exists; reports are never overwritten"}, code=2)
    project = data.get("project", "")
    pn = data.get("platform_name") or ""
    data["title"] = (f"{project} – {data['plan_name']} {pn} Test Report" if project
                     else f"{data['plan_name']} {pn} Test Report").replace("  ", " ")
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    work = out + ".building.xlsx"  # built aside, moved into place only when complete
    try:
        shutil.copyfile(TEMPLATE, work)
        wb = openpyxl.load_workbook(work)
        for ws in list(wb.worksheets):
            if ws.title.strip().lower() == "report templates":
                wb.remove(ws)
        if "Report Template" in wb.sheetnames:
            wb["Report Template"].title = REPORT_SHEET
        fill_document_control(wb["Document Control"], data)
        fill_version_history(wb["Version History"], data)
        ranges = fill_report_sheet(wb[REPORT_SHEET], data)
        xlsx_layout.tidy_workbook(wb)  # widths, wrapping, alignment, row heights - see xlsx_layout.py
        # the COUNTIF formulas must give the same numbers as the MD report
        ws = wb[REPORT_SHEET]
        seen = {s: 0 for s in STATUS_LIST}
        for rr in range(FIRST_BAND_ROW, ws.max_row + 1):
            v = ws.cell(rr, 6).value
            if v in seen:
                seen[v] += 1
        c = data["counts"]
        expected = {"PASS": c["passed"], "FAIL": c["failed"], "IN PROGRESS": c["in_progress"],
                    "UNEXECUTED": c["unexecuted"], "BLOCKED": c["blocked"]}
        if seen != expected:
            dump({"ok": False, "error": "Status cells do not match the report counts", "sheet": seen,
                  "report": expected}, code=1)
        wb.calculation.fullCalcOnLoad = True
        wb.active = wb.sheetnames.index(REPORT_SHEET)
        wb.save(work)
        recalculated = False if a.no_libreoffice else recalc_with_libreoffice(work)
        os.replace(work, out)
    finally:
        if os.path.exists(work):  # never leave a half-built file behind
            os.remove(work)
    dump({"ok": True, "xlsx": rel(out, root), "sheets": wb.sheetnames, "status_ranges": ranges,
          "recalculation": "LibreOffice" if recalculated else "on open (fullCalcOnLoad)",
          "counts": data["counts"]})


if __name__ == "__main__":
    main()
