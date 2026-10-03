#!/usr/bin/env python3
"""Write execution results into the test case MD, then the XLSX (Phase 6 of execute-mobile-test-cases).
Matching is by Test Case ID (same behaviour as the web skill).

  python update_results.py "<test cases .md>" --results <results.json> [--xlsx "<test cases .xlsx>"]
                           [--bug-index bugs/bug-index.md] [--no-backup]

results.json:
  {"results": [
     {"id": "TC-UI-NEG-014",
      "actual": "No error message was displayed after clicking Subscribe.",
      "status": "FAIL",                         # PASS | FAIL | BLOCKED | IN PROGRESS | PENDING
      "executed_on": "03-10-2026 14:35:10 IST",
      "defects": "BUG-003",
      "comments": "",
      "corrections": {"Expected Result": "..."} # optional, only for a corrected test case
     }]}

Only the keys present are written (an empty string clears the cell). Columns written: Actual Result,
Execution Status, Executed On, Execution Defects, Comments (+ the wording columns named in `corrections`).
Created By, Reviewed By and Executed By are never touched. [NEW]/[UPDATED] tags in Comments are kept.

MD: a copy of the previous version goes to `Test Cases/.history/` first, the table cells are replaced and the
"Latest execution summary" section is refreshed; the file is written in one step.
Jira links: bugs recorded in bugs/jira-sync.json (file-bugs-to-jira) are written to Execution Defects as
"[PROJ-45](<url>) (BUG-003)" in the MD and "PROJ-45 (BUG-003)" with a cell hyperlink in the XLSX (jira_refs.py), so a
results update never loses a Jira link.
XLSX (only if given and present): columns found by header name, same rows updated, styles, dropdown,
conditional colours, merged flow bands, frozen header, Document Control and Version History untouched.
"""
import argparse
import os
import re
import shutil
import sys
from collections import Counter, OrderedDict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import (COLUMNS, PROTECTED, STATUSES, cell_out, dump, find_root, ist_cell, ist_file,  # noqa: E402
                     is_obsolete, load_json, norm_header, parse_testcase_md, read_tables, rel, utf8_stdio,
                     write_text_atomic)
import jira_refs  # noqa: E402
import xlsx_layout  # noqa: E402

FIELD_TO_COL = OrderedDict([("actual", "Actual Result"), ("status", "Execution Status"),
                            ("executed_on", "Executed On"), ("defects", "Execution Defects"),
                            ("comments", "Comments")])
CORRECTABLE = ["Test Description", "Preconditions", "Test Steps", "Test Data", "Expected Result"]
EXEC_ON_RE = re.compile(r"^\d{2}-\d{2}-\d{4} \d{2}:\d{2}:\d{2} IST$")
TAG_RE = re.compile(r"^\s*(\[(?:NEW|UPDATED)\])\s*", re.I)
SUMMARY_H2 = "## Latest execution summary"


def validate(results, known_ids):
    errors = []
    for r in results:
        cid = r.get("id", "")
        if cid not in known_ids:
            errors.append(f"{cid or '<no id>'}: not found in the test case file")
        st = r.get("status")
        if st is not None and st not in STATUSES:
            errors.append(f"{cid}: status '{st}' is not one of {STATUSES}")
        eo = r.get("executed_on")
        if eo and not EXEC_ON_RE.match(eo):
            errors.append(f"{cid}: executed_on '{eo}' must be dd-mm-yyyy hh:mm:ss IST")
        for col in (r.get("corrections") or {}):
            if col not in CORRECTABLE:
                errors.append(f"{cid}: '{col}' cannot be corrected (allowed: {CORRECTABLE})")
        if r.get("corrections") and "Test case corrected" not in (r.get("comments") or ""):
            errors.append(f"{cid}: a corrected case needs a Comments note 'Test case corrected: …'")
    return errors


def merged_comment(old, new):
    m = TAG_RE.match(old or "")
    if m and not TAG_RE.match(new or ""):
        return (m.group(1).upper() + " " + (new or "")).strip()
    return new


def new_values(case, r, sync=None):
    vals = {}
    for key, col in FIELD_TO_COL.items():
        if key in r:
            v = r[key] if r[key] is not None else ""
            if col == "Comments":
                v = merged_comment(case.get("Comments", ""), v)
            if col == "Execution Defects":
                v = jira_refs.render_md(v, sync or {})
            vals[col] = v
    for col, v in (r.get("corrections") or {}).items():
        vals[col] = v
    return vals


# ---------------------------------------------------------------- bug index

def read_bug_index(path):
    bugs = OrderedDict()
    if not path or not os.path.exists(path):
        return bugs
    with open(path, encoding="utf-8") as f:
        lines = f.read().split("\n")
    for header, rows, _, _ in read_tables(lines, 0, len(lines)):
        keys = [norm_header(h) for h in header]
        if "bug id" not in keys:
            continue
        for _, cells in rows:
            row = {keys[i]: (cells[i] if i < len(cells) else "") for i in range(len(keys))}
            bid = re.sub(r"[\[\]]", "", row.get("bug id", "")).strip()
            bid = re.match(r"(BUG-\d+)", bid).group(1) if re.match(r"(BUG-\d+)", bid) else bid
            if bid:
                bugs[bid] = row
    return bugs


# ---------------------------------------------------------------- MD

def summary_section(doc_cases_by_flow, bugs, referenced, sync=None):
    sync = sync or {}
    lines = [SUMMARY_H2, "", f"Last updated: {ist_cell()}", "",
             "| Flow | Total | PASS | FAIL | BLOCKED | IN PROGRESS | PENDING |",
             "|---|---|---|---|---|---|---|"]
    tot = Counter()
    for name, cases in doc_cases_by_flow:
        c = Counter((x.get("Execution Status") or "PENDING").strip().upper() for x in cases)
        n = len(cases)
        tot["n"] += n
        for s in STATUSES:
            tot[s] += c.get(s, 0)
        lines.append(f"| {cell_out(name)} | {n} | {c.get('PASS', 0)} | {c.get('FAIL', 0)} | {c.get('BLOCKED', 0)} | "
                     f"{c.get('IN PROGRESS', 0)} | {c.get('PENDING', 0)} |")
    lines.append(f"| **Total** | **{tot['n']}** | **{tot['PASS']}** | **{tot['FAIL']}** | **{tot['BLOCKED']}** | "
                 f"**{tot['IN PROGRESS']}** | **{tot['PENDING']}** |")
    lines += ["", "Obsolete cases are not counted.", ""]
    shown = [b for b in bugs if b in referenced]
    if shown:
        jira = any(b in sync for b in shown)
        lines += ["| Bug ID | Title | Severity | Status | Test case(s) |" + (" Jira |" if jira else ""),
                  "|---|---|---|---|---|" + ("---|" if jira else "")]
        for b in shown:
            row = bugs[b]
            link = (f" [{sync[b]['jiraKey']}]({sync[b]['url']}) |" if b in sync else " Not filed |") if jira else ""
            lines.append(f"| {b} | {row.get('title', '')} | {row.get('severity', '')} | {row.get('status', '')} | "
                         f"{row.get('test case', row.get('test cases', ''))} |" + link)
    else:
        lines.append("No defects raised.")
    lines.append("")
    return lines


def update_md(md, results, bug_index, backup, root, sync=None):
    doc = parse_testcase_md(md)
    lines = list(doc["lines"])
    by_id = {r["id"]: r for r in results}
    changed = []
    for f in doc["flows"]:
        cm = f["colmap"]
        for case in f["cases"]:
            r = by_id.get(case.get("Test Case ID"))
            if not r:
                continue
            vals = new_values(case, r, sync)
            cells = list(case["raw_cells"]) + [""] * max(0, len(cm) - len(case["raw_cells"]))
            for col, v in vals.items():
                if col in PROTECTED or col not in cm:
                    continue
                cells[cm[col]] = cell_out(v)
                case[col] = v
            lines[case["line"]] = "| " + " | ".join(cells) + " |"
            changed.append(case["Test Case ID"])
    # refresh summary section
    referenced = set()
    flows_cases = []
    for f in doc["flows"]:
        live = [c for c in f["cases"] if not is_obsolete(c)]
        flows_cases.append((f["name"], live))
        for c in live:
            referenced.update(re.findall(r"BUG-\d+", c.get("Execution Defects", "") or ""))
    section = summary_section(flows_cases, read_bug_index(bug_index), referenced, sync)
    start = next((i for i, l in enumerate(lines) if l.strip() == SUMMARY_H2), None)
    if start is not None:
        end = next((i for i in range(start + 1, len(lines)) if lines[i].startswith("## ")), len(lines))
        lines[start:end] = section
    else:
        obs = next((i for i, l in enumerate(lines) if l.strip().lower().startswith("## observations")), None)
        at = obs if obs is not None else len(lines)
        if at == len(lines) and lines and lines[-1].strip():
            section = [""] + section
        lines[at:at] = section
    backup_path = None
    if backup:
        hist = os.path.join(os.path.dirname(os.path.abspath(md)), ".history")
        os.makedirs(hist, exist_ok=True)
        stem, ext = os.path.splitext(os.path.basename(md))
        backup_path = os.path.join(hist, f"{stem}__{ist_file()}{ext}")
        shutil.copy2(md, backup_path)
    write_text_atomic(md, "\n".join(lines))
    return changed, backup_path


# ---------------------------------------------------------------- XLSX

def find_case_sheet(wb):
    for ws in wb.worksheets:
        for r in range(1, 8):
            for c in range(1, min(ws.max_column, 40) + 1):
                if norm_header(ws.cell(r, c).value) == "test case id":
                    return ws, r
    return None, None


def set_defects_cell(cell, md_text, sync):
    """Execution Defects in the XLSX: plain 'PROJ-45 (BUG-003)' with the issue as the cell's hyperlink."""
    value, link = jira_refs.render_plain(md_text, sync)
    cell.value = value if value else None
    cell.hyperlink = link if link else None


def update_xlsx(xlsx, md_after, results, backup, sync=None):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    ws, hrow = find_case_sheet(wb)
    if ws is None:
        return {"error": "No sheet with a 'Test Case ID' header found"}
    canon = {norm_header(c): c for c in COLUMNS}
    cols = {}
    for c in range(1, ws.max_column + 1):
        k = norm_header(ws.cell(hrow, c).value)
        if k in canon:
            cols[canon[k]] = c
    # values come from the MD just written, so MD and XLSX always agree
    doc = parse_testcase_md(md_after)
    md_cases = {c["Test Case ID"]: c for f in doc["flows"] for c in f["cases"]}
    by_id = {r["id"]: r for r in results}
    id_col = cols["Test Case ID"]
    changed, missing = [], set(by_id)
    for row in range(hrow + 1, ws.max_row + 1):
        cid = ws.cell(row, id_col).value
        cid = str(cid).strip() if cid is not None else ""
        if cid not in by_id:
            continue
        missing.discard(cid)
        r, case = by_id[cid], md_cases.get(cid, {})
        touch = [col for key, col in FIELD_TO_COL.items() if key in r] + list((r.get("corrections") or {}).keys())
        for col in touch:
            if col in PROTECTED or col not in cols:
                continue
            v = case.get(col, "")
            if col == "Execution Defects":
                set_defects_cell(ws.cell(row, cols[col]), v, sync or {})
                continue
            ws.cell(row, cols[col]).value = v if v != "" else None
        changed.append(cid)
    xlsx_layout.tidy_workbook(wb)  # new Actual Result / Comments text must wrap and fit its row
    backup_path = None
    if backup:
        hist = os.path.join(os.path.dirname(os.path.abspath(xlsx)), ".history")
        os.makedirs(hist, exist_ok=True)
        stem, ext = os.path.splitext(os.path.basename(xlsx))
        backup_path = os.path.join(hist, f"{stem}__{ist_file()}{ext}")
        shutil.copy2(xlsx, backup_path)
    tmp = xlsx + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, xlsx)
    return {"sheet": ws.title, "updated": len(changed), "not_found_in_xlsx": sorted(missing),
            "backup": backup_path}


def main():
    utf8_stdio()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("md")
    ap.add_argument("--results", required=True)
    ap.add_argument("--xlsx")
    ap.add_argument("--bug-index", default="bugs/bug-index.md")
    ap.add_argument("--jira-sync", default="bugs/jira-sync.json")
    ap.add_argument("--no-backup", action="store_true")
    a = ap.parse_args()
    root = find_root()
    if not os.path.isfile(a.md):
        dump({"ok": False, "error": f"File not found: {a.md}"}, code=2)
    data = load_json(a.results, {}) or {}
    results = data.get("results", data if isinstance(data, list) else [])
    doc = parse_testcase_md(a.md)
    known = {c["Test Case ID"] for f in doc["flows"] for c in f["cases"]}
    errors = validate(results, known)
    if errors:
        dump({"ok": False, "errors": errors}, code=1)
    sync = jira_refs.load_sync(a.jira_sync)
    changed, md_backup = update_md(a.md, results, a.bug_index, not a.no_backup, root, sync)
    out = {"ok": True, "md": rel(a.md, root), "md_updated": len(changed),
           "md_backup": rel(md_backup, root) if md_backup else None}
    if a.xlsx:
        if os.path.exists(a.xlsx):
            x = update_xlsx(a.xlsx, a.md, results, not a.no_backup, sync)
            if x.get("backup"):
                x["backup"] = rel(x["backup"], root)
            out["xlsx"] = {"path": rel(a.xlsx, root), **x}
            if x.get("error"):
                out["ok"] = False
        else:
            out["xlsx"] = {"path": a.xlsx, "skipped": "file does not exist"}
    dump(out, code=0 if out["ok"] else 1)


if __name__ == "__main__":
    main()
