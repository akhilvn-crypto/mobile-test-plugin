#!/usr/bin/env python3
"""Build the XLSX from a build-mobile-test-cases Markdown file, on the user's template.

  python md_to_xlsx.py <file.md> [--out <file.xlsx>] [--template <Testcases_Template.xlsx>]

- Copies assets/Testcases_Template.xlsx, deletes the Test Suite sheet.
- Document Control: template rows kept, sample values cleared; only Title, Description and
  Privacy Classification are filled. Description carries the platform, device, OS and app build:
  'Test cases for <plan>, Android, Pixel 7, Android 14, app 2.4.1 (build 123)'.
- Version History: template columns kept, sample rows cleared, rows taken from the MD's Version History.
- Testcases: 15 columns (Actual Result inserted after Expected Result, so Execution Status moves H -> I),
  one blue band row per flow (4A86E8, merged A:C), frozen header, wrapped/top-aligned Calibri 11,
  Execution Status dropdown + the template's conditional colours rebuilt on column I.
- Layout: xlsx_layout.tidy_workbook() runs last on every sheet - readable column widths, headers wrapped
  and centred, data wrapped/top/left (status, people, date and defect columns centred), every row tall
  enough for its text so nothing overflows or is clipped.
- Never overwrites an existing file.
"""
import argparse
import copy
import os
import re
import sys

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lint_testcases import COLUMNS, STATUSES, parse_md  # noqa: E402
import xlsx_layout  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "..", "assets", "Testcases_Template.xlsx")
BAND_FILL = "FF4A86E8"

# new column (1-based) -> template column it takes header style / width from
SRC_COL = {1: 1, 2: 2, 3: 3, 4: 4, 5: 5, 6: 6, 7: 7, 8: 7, 9: 8, 10: 9, 11: 10, 12: 11, 13: 12, 14: 13, 15: 14}
HEADERS = ["Test Case ID", "Test Case Title", "Test Description", "Preconditions", "Test Steps", "Test Data",
           "Expected Result", "Actual Result", "Execution\nStatus", "Comments", "Created\nBy", "Reviewed By",
           "Executed\nBy", "Executed On\n(dd-mm-yyyy hh:mm:ss IST)", "Execution\nDefects"]
WIDTH_OVERRIDE = {1: 18.0, 8: 30.0, 9: 13.0, 14: 22.0}  # A widened for new IDs; Actual Result; status; Executed On
STATUS_COL = "I"


def cell_text(v):
    v = re.sub(r"<br\s*/?>", "\n", v or "", flags=re.I)
    return v.strip() or None


def copy_style(src, dst):
    dst.font = copy.copy(src.font)
    dst.fill = copy.copy(src.fill)
    dst.border = copy.copy(src.border)
    dst.alignment = copy.copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy.copy(src.protection)


def ranges(rows):
    """[3,4,5,7] -> 'I3:I5 I7'"""
    out, start, prev = [], None, None
    for r in sorted(rows):
        if start is None:
            start = prev = r
        elif r == prev + 1:
            prev = r
        else:
            out.append((start, prev))
            start = prev = r
    if start is not None:
        out.append((start, prev))
    return " ".join(f"{STATUS_COL}{a}" if a == b else f"{STATUS_COL}{a}:{STATUS_COL}{b}" for a, b in out)


def fill_document_control(ws, title, description):
    labels = {}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip() in (
                    "Title", "Document ID", "Description", "Author", "Review Date", "Reviewed by",
                    "Approved date", "Approved By", "Privacy Classification"):
                labels[c.value.strip()] = c
    for name, lc in labels.items():
        ws.cell(row=lc.row, column=lc.column + 1).value = None
    values = {"Title": title, "Description": description, "Privacy Classification": "Confidential"}
    for name, val in values.items():
        if name in labels:
            lc = labels[name]
            ws.cell(row=lc.row, column=lc.column + 1).value = val


def fill_version_history(ws, rows):
    hdr = None
    for row in ws.iter_rows():
        for c in row:
            if c.value == "Version No":
                hdr = c
                break
        if hdr:
            break
    if hdr is None:
        raise SystemExit("Version History sheet: 'Version No' header not found in template")
    cols = []
    c = hdr.column
    while ws.cell(row=hdr.row, column=c).value:
        cols.append((c, ws.cell(row=hdr.row, column=c).value))
        c += 1
    sample = {col: copy.copy(ws.cell(row=hdr.row + 1, column=col)._style) for col, _ in cols}
    for r in range(hdr.row + 1, ws.max_row + 1):
        for col, _ in cols:
            ws.cell(row=r, column=col).value = None
    for i, vh in enumerate(rows):
        r = hdr.row + 1 + i
        for col, name in cols:
            cell = ws.cell(row=r, column=col)
            cell._style = copy.copy(sample[col])
            v = vh.get(name, "")
            cell.value = v if name in ("Version No", "Created On", "Comments") and v else None
            if name == "Comments":
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_testcases(ws, doc):
    hdr_styles = {c: copy.copy(ws.cell(row=1, column=c)._style) for c in range(1, 15)}
    band_style = copy.copy(ws["A2"]._style)
    data_style = copy.copy(ws["B3"]._style)
    widths = {c: ws.column_dimensions[chr(64 + c)].width for c in range(1, 15)}
    cf_rules = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            cf_rules.append(rule)
    row1_ht = ws.row_dimensions[1].height

    # clear everything below the header; rebuild merged cells, validation and CF explicitly
    for mr in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mr))
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r in list(ws.row_dimensions):
        if r != 1:
            del ws.row_dimensions[r]
    ws.data_validations.dataValidation = []
    ws.conditional_formatting = ConditionalFormattingList()
    for c in range(1, ws.max_column + 1):
        ws.cell(row=1, column=c).value = None

    # header (15 columns)
    for c, text in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=text)
        cell._style = copy.copy(hdr_styles[SRC_COL[c]])
    ws.row_dimensions[1].height = max(row1_ht or 15.75, 45)
    letters = "ABCDEFGHIJKLMNO"
    for c in range(1, 16):
        ws.column_dimensions[letters[c - 1]].width = WIDTH_OVERRIDE.get(c, widths[SRC_COL[c]])

    r = 2
    status_rows = []
    for flow in doc["flows"]:
        band = ws.cell(row=r, column=1, value=flow["name"])
        band._style = copy.copy(band_style)
        band.fill = PatternFill("solid", fgColor=BAND_FILL)
        for c in (2, 3):
            ws.cell(row=r, column=c)._style = copy.copy(band_style)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        r += 1
        for case in flow["cases"]:
            for c, col in enumerate(COLUMNS, start=1):
                v = cell_text(case.get(col, ""))
                if col == "Execution Status":
                    v = v if v in STATUSES else "PENDING"
                cell = ws.cell(row=r, column=c, value=v)
                cell._style = copy.copy(data_style)
                cell.font = Font(name="Calibri", size=11)
            status_rows.append(r)
            r += 1

    last = max(r - 1, 1000)
    if status_rows:
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
        for rng in ranges(status_rows).split():
            dv.add(rng)
        ws.add_data_validation(dv)
    for rule in cf_rules:
        ws.conditional_formatting.add(f"{STATUS_COL}1:{STATUS_COL}{last}", rule)
    ws.freeze_panes = "D2"



def description(plan_name, header):
    """Test cases for <plan>, Android, Pixel 7, Android 14, app 2.4.1 (build 123) - from the MD header rows."""
    parts = [f"Test cases for {plan_name}"]
    for key in ("Platform", "Device", "OS version"):
        v = re.sub(r"\s*\(real device\)\s*$", "", (header.get(key) or "").strip())
        if v:
            parts.append(v)
    app = (header.get("App version and build") or "").strip()
    if app:
        parts.append(app if app.lower().startswith("app") else f"app {app}")
    return ", ".join(parts)


def main():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("md")
    ap.add_argument("--out")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    a = ap.parse_args()

    out = a.out or os.path.splitext(a.md)[0] + ".xlsx"
    if os.path.exists(out):
        sys.exit(f"refusing to overwrite existing file: {out}")
    if not os.path.isfile(a.template):
        sys.exit(f"template not found: {a.template}")
    doc = parse_md(a.md)
    if not doc["flows"]:
        sys.exit("no flow sections found in the MD")
    fm = doc["front_matter"] or {}
    plan_name = doc["header"].get("Plan name") or (doc["h1"] or "").split(" – ")[0]
    title = fm.get("title") or f"{plan_name} Test Cases"

    wb = load_workbook(a.template)
    for name in ("Document Control", "Version History", "Testcases"):
        if name not in wb.sheetnames:
            sys.exit(f"template sheet missing: {name}")
    if "Test Suite" in wb.sheetnames:
        del wb["Test Suite"]
    fill_document_control(wb["Document Control"], title, description(plan_name, doc["header"]))
    vh = doc["version_history"] or [{"Version No": fm.get("version", "V1.0"), "Comments": "Initial generation"}]
    fill_version_history(wb["Version History"], vh)
    build_testcases(wb["Testcases"], doc)
    xlsx_layout.tidy_workbook(wb)  # widths, wrapping, alignment, row heights - see xlsx_layout.py
    wb.active = wb.sheetnames.index("Testcases")
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Testcases"
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    n = sum(len(f["cases"]) for f in doc["flows"])
    print(f"wrote {out} ({len(doc['flows'])} flow(s), {n} case(s), sheets: {', '.join(wb.sheetnames)})")


if __name__ == "__main__":
    main()
