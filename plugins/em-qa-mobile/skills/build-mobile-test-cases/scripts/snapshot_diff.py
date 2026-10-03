#!/usr/bin/env python3
"""Exploration snapshot utilities for build-mobile-test-cases (one platform per suite).

Subcommands
  plan             Split the test plan into plan name, common part, flows and environment hints (JSON).
  compare          Compare docs / plan / app build hash / screen fingerprints with manifest.json for ONE platform
                   and print the decision: full | partial | check-drift | stop (JSON).
  fingerprint      Print (or --write into the files) the structure fingerprint of screen JSON files.
  archive          Move exploration files into history/<IST-timestamp>/ (never deletes).
  scrub            Mask tokens, cookies and credentials in a device log, JSON or text file (alias: scrub-log).
  update-manifest  Write the platform's part of manifest.json after a run (hashes, app build, flows, screens, runs[]).
  update-app       Record a new app build whose light check found no changed screen (no output file is written).

Layout: Exploration/<plan-key>/{manifest.json, project-understanding.md, <platform>/<flow-key>/..., runs/<ts>/}.
Run `snapshot_diff.py <subcommand> -h` for options. See references/incremental-mode.md.
"""
import argparse
import glob
import hashlib
import json
import os
import re
import shutil
import sys
import zipfile
from datetime import datetime, timedelta, timezone

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import check_prerequisites as prereq  # noqa: E402

IST = timezone(timedelta(hours=5, minutes=30), "IST")
MASK = "********"
PLATFORMS = ("android", "ios")
FLOW_PARENT_NAMES = {"flows", "test flows", "user flows", "scenarios", "test scenarios", "flows to test"}
COMMON_SECTIONS = {
    "scope", "in scope", "out of scope", "out-of-scope", "environment", "environments", "url", "urls",
    "credentials", "roles", "users", "user roles", "test data", "supporting data", "data", "assumptions",
    "notes", "overview", "introduction", "objective", "objectives", "schedule", "risks", "references",
    "allowed actions", "general", "summary", "contacts",
    # mobile
    "platform", "platforms", "app", "apps", "application", "app build", "app builds", "build", "builds",
    "device", "devices", "test device", "test devices", "appium", "appium server", "permissions",
    "permissions allowed", "test environment", "mobile", "app under test",
}
SKIP_DIRS = {"history", "runs"} | set(PLATFORMS)


# ---------------------------------------------------------------- helpers

def ist_file_ts():
    return datetime.now(IST).strftime("%Y-%m-%d_%H-%M-%S") + "-IST"


def ist_cell_ts(dt=None):
    return (dt or datetime.now(IST)).strftime("%d-%m-%Y %H:%M:%S") + " IST"


def sha(data):
    if isinstance(data, str):
        data = data.encode("utf-8")
    return hashlib.sha256(data).hexdigest()


def slugify(text):
    text = re.sub(r"[^a-z0-9]+", "-", text.lower())
    return text.strip("-") or "item"


def normalise(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [l.rstrip() for l in text.split("\n")]
    out, blank = [], False
    for l in lines:
        if not l:
            if not blank:
                out.append("")
            blank = True
        else:
            out.append(l)
            blank = False
    return "\n".join(out).strip() + "\n"


def file_hash(path):
    with open(path, "rb") as f:
        return sha(f.read())


def list_docs(project_dir):
    docs = {}
    for dp, dns, fs in os.walk(project_dir):
        dns[:] = [d for d in dns if not d.startswith(".")]
        for f in fs:
            if f.startswith(("~$", ".")):
                continue
            p = os.path.join(dp, f)
            docs[os.path.relpath(p, project_dir).replace("\\", "/")] = file_hash(p)
    return dict(sorted(docs.items()))


def load_json(path, default=None):
    if not path or not os.path.isfile(path):
        return default
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def dump(obj):
    json.dump(obj, sys.stdout, indent=2, ensure_ascii=False)
    print()


# ---------------------------------------------------------------- plan reading

def docx_to_markdown(path):
    """Word -> pseudo-Markdown using heading styles. python-docx if present, else raw XML."""
    try:
        import docx  # type: ignore
        d = docx.Document(path)
        lines = []
        for p in d.paragraphs:
            style = (p.style.name or "").lower() if p.style is not None else ""
            m = re.match(r"heading (\d)", style)
            if style == "title":
                lines.append("# " + p.text.strip())
            elif m:
                lines.append("#" * int(m.group(1)) + " " + p.text.strip())
            else:
                lines.append(p.text)
        for t in d.tables:
            for r in t.rows:
                lines.append("| " + " | ".join(c.text.strip() for c in r.cells) + " |")
        return "\n".join(lines)
    except ImportError:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
        lines = []
        for para in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
            text = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", para, re.S))
            m = re.search(r'<w:pStyle w:val="(?:Heading|heading)(\d)"', para)
            if re.search(r'<w:pStyle w:val="Title"', para):
                lines.append("# " + text)
            elif m:
                lines.append("#" * int(m.group(1)) + " " + text)
            else:
                lines.append(text)
        return "\n".join(lines)


def read_plan_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return docx_to_markdown(path), True
    if ext == ".pdf":
        try:
            import pypdf  # type: ignore
            r = pypdf.PdfReader(path)
            return "\n".join((pg.extract_text() or "") for pg in r.pages), False
        except ImportError:
            return None, False
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read(), True


def strip_flow_prefix(name):
    n = re.sub(r"^\s*flow\s*\d*\s*[:.\-–—)]?\s*", "", name, flags=re.I)
    n = re.sub(r"^\s*\d+(\.\d+)*[.)]?\s+", "", n)
    return n.strip() or name.strip()


def parse_plan(path):
    raw, splittable = read_plan_text(path)
    base = os.path.splitext(os.path.basename(path))[0]
    if raw is None:  # unreadable PDF
        return {"plan_name": base, "plan_hash": file_hash(path), "common_hash": file_hash(path),
                "splittable": False, "flows": [], "warning": "PDF plan could not be read by the script; "
                "install pypdf or provide the plan as .md/.docx. Every flow counts as changed when the file changes."}
    text = normalise(raw)
    title = None
    body = text
    fm = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if fm:
        t = re.search(r"^title:\s*[\"']?(.*?)[\"']?\s*$", fm.group(1), re.M)
        if t:
            title = t.group(1).strip()
        body = text[fm.end():]

    lines = body.split("\n")
    heads, fence = [], False
    for i, l in enumerate(lines):
        if l.lstrip().startswith(("```", "~~~")):
            fence = not fence
            continue
        m = None if fence else re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", l)
        if m:
            heads.append({"i": i, "level": len(m.group(1)), "text": m.group(2).strip()})

    if not title:
        h1 = next((h for h in heads if h["level"] == 1), None)
        title = h1["text"] if h1 else next((l.strip() for l in lines if l.strip()), base)

    def section_end(idx):
        h = heads[idx]
        for j in range(idx + 1, len(heads)):
            if heads[j]["level"] <= h["level"]:
                return heads[j]["i"]
        return len(lines)

    flow_idx = []
    parent = next((k for k, h in enumerate(heads) if h["text"].lower().rstrip(":") in FLOW_PARENT_NAMES), None)
    if parent is not None:
        end = section_end(parent)
        kids = [k for k, h in enumerate(heads) if heads[parent]["i"] < h["i"] < end]
        if kids:
            lvl = min(heads[k]["level"] for k in kids)
            flow_idx = [k for k in kids if heads[k]["level"] == lvl]
    if not flow_idx:
        flow_idx = [k for k, h in enumerate(heads) if re.match(r"^flow\b", h["text"], re.I)]
    if not flow_idx:
        flow_idx = [k for k, h in enumerate(heads) if h["level"] == 2
                    and re.sub(r"[^a-z -]", "", h["text"].lower()).strip() not in COMMON_SECTIONS]

    flows, used, taken = [], set(), set()
    for k in flow_idx:
        start, end = heads[k]["i"], section_end(k)
        sec = "\n".join(lines[start:end])
        name = strip_flow_prefix(heads[k]["text"])
        key = slugify(name)
        base_key, n = key, 2
        while key in used:
            key, n = f"{base_key}-{n}", n + 1
        used.add(key)
        taken.update(range(start, end))
        flows.append({"flow_key": key, "name": name, "heading": heads[k]["text"],
                      "section_hash": sha(normalise(sec)), "line": start + 1})
    common = "\n".join(l for i, l in enumerate(lines) if i not in taken)
    return {"plan_name": title, "plan_hash": sha(text), "common_hash": sha(normalise(common)),
            "splittable": splittable, "flows": flows}


def environment_hints(path):
    """Best-effort hints from the plan text: platforms, app file / id per platform, device, Appium address.
    The skill reads the plan itself and decides; these hints only save time."""
    raw, _ = read_plan_text(path)
    text = raw or ""
    low = text.lower()
    platforms = [p for p in PLATFORMS if re.search(r"\b" + ("android" if p == "android" else r"ios|iphone|ipad") + r"\b", low)]
    apps = {}
    for m in re.finditer(r"([\w./\\:~ -]+?\.(?:apk|aab|ipa))\b", text, re.I):
        f = m.group(1).strip().strip("`\"' :")
        f = re.sub(r"^.*?:\s+", "", f) if re.match(r"^[A-Za-z ]+:\s", f) else f
        apps.setdefault("android" if f.lower().endswith((".apk", ".aab")) else "ios", f)
    ids = re.findall(r"(?:package(?: name)?|bundle(?: id| identifier)?|app id)\s*[:=]\s*`?([A-Za-z][\w]*(?:\.[\w]+)+)`?",
                     text, re.I)
    devices = re.findall(r"(?:device|phone|handset)\s*[:=]\s*([^\n|]+)", text, re.I)
    url = re.search(r"(https?://(?:127\.0\.0\.1|localhost|[\w.-]+):\d+(?:/wd/hub)?)", text)
    return {"platforms": platforms, "app_files": apps, "app_ids": ids[:4], "devices": [d.strip() for d in devices[:4]],
            "appium_url": url.group(1) if url else None}


# ---------------------------------------------------------------- fingerprints

def _norm_str(s):
    s = re.sub(r"\s+", " ", str(s)).strip().lower()
    return re.sub(r"\d+", "#", s)


def fingerprint(screen):
    """Structure fingerprint of a screen JSON: title, controls (label + type + has label/key), input fields
    (label + type + required). Enabled state, sizes and messages are left out (they change at run time)."""
    controls = []
    for c in screen.get("controls", []) or []:
        if isinstance(c, dict):
            controls.append("|".join([_norm_str(c.get("label", "")), _norm_str(c.get("type", "")),
                                      _norm_str(c.get("key", "")),
                                      "labelled" if c.get("has_label") or c.get("has_key") else "unlabelled"]))
        else:
            controls.append(_norm_str(c))
    fields = []
    for f in screen.get("input_fields", []) or []:
        if isinstance(f, dict):
            fields.append("|".join(_norm_str(f.get(k, "")) for k in ("label", "type", "required")))
        else:
            fields.append(_norm_str(f))
    canon = {"title": _norm_str(screen.get("title", "")), "controls": sorted(controls), "input_fields": sorted(fields)}
    return sha(json.dumps(canon, sort_keys=True))


def flow_screens(flow_dir):
    screens = []
    for p in sorted(glob.glob(os.path.join(flow_dir, "screens", "*.json"))):
        try:
            s = load_json(p, {})
        except (json.JSONDecodeError, OSError):
            continue
        captured = s.get("captured_ist") or ist_cell_ts(datetime.fromtimestamp(os.path.getmtime(p), IST))
        entry = {"screen": s.get("screen") or os.path.splitext(os.path.basename(p))[0],
                 "reached_by": s.get("reached_by", ""), "fingerprint": fingerprint(s), "last_captured_ist": captured}
        if s.get("updated_by"):  # refreshed by execute-mobile-test-cases: a valid capture, keep the marker
            entry["updated_by"] = s["updated_by"]
        screens.append(entry)
    return screens


# ---------------------------------------------------------------- manifest access

def load_manifest(plan_dir):
    return load_json(os.path.join(plan_dir, "manifest.json"), None)


def platform_part(manifest, platform):
    return ((manifest or {}).get("platforms") or {}).get(platform)


def app_info(app, platform, serial):
    if not app:
        return {}
    info = prereq.inspect_app(app, platform, serial, scan=False)
    info["fingerprint"] = prereq.app_fingerprint(info)
    return info


# ---------------------------------------------------------------- plan / compare

def cmd_plan(a):
    out = parse_plan(a.test_plan)
    out["environment"] = environment_hints(a.test_plan)
    dump(out)


def cmd_compare(a):
    plan_dir = a.plan_dir
    manifest = load_manifest(plan_dir)
    part = platform_part(manifest, a.platform)
    plan = parse_plan(a.test_plan)
    docs = list_docs(a.project_dir)
    app = app_info(a.app, a.platform, a.serial)
    res = {"plan_key": os.path.basename(os.path.normpath(plan_dir)), "plan_name": plan["plan_name"],
           "platform": a.platform, "platform_dir": os.path.join(plan_dir, a.platform).replace("\\", "/"),
           "decision": None, "reason": "", "splittable": plan["splittable"], "needs_device": True,
           "app": {"given": a.app, "path": app.get("path"), "package": app.get("package"), "version": app.get("version"),
                   "build": app.get("build"), "hash": app.get("fingerprint"), "previous_hash": None, "changed": None},
           "hashes": {"plan_hash": plan["plan_hash"], "common_hash": plan["common_hash"], "doc_hashes": docs,
                      "flows": {f["flow_key"]: {"name": f["name"], "section_hash": f["section_hash"]} for f in plan["flows"]}},
           "docs": {"changed": [], "added": [], "removed": []},
           "common_changed": False,
           "flows": {"new": [], "changed": [], "removed": [], "unchanged": []},
           "screens_drifted": {}, "actions": []}
    if plan.get("warning"):
        res["warning"] = plan["warning"]
    if a.app and not app.get("fingerprint"):
        res["app_warning"] = ("The app build fingerprint could not be computed (file missing, or app not installed "
                              "on the device); every unchanged flow gets the light check.")

    if part is None or a.force_full:
        res["decision"] = "full"
        res["reason"] = "--force-full set" if a.force_full else f"no {a.platform} exploration in manifest.json (first run)"
        res["flows"]["new"] = [f["flow_key"] for f in plan["flows"]]
        res["actions"] = [{"flow_key": f["flow_key"], "action": "full-explore"} for f in plan["flows"]]
        if part and a.force_full:
            res["previous_version"] = part.get("version")
            res["previous_output"] = part.get("last_output_file")
        return dump(res)

    res["previous_version"] = part.get("version")
    res["previous_output"] = part.get("last_output_file")
    prev_hash = part.get("app_hash")
    res["app"]["previous_hash"] = prev_hash
    app_changed = (not app.get("fingerprint")) or prev_hash != app.get("fingerprint")
    res["app"]["changed"] = app_changed
    old_docs = part.get("doc_hashes", {})
    res["docs"]["changed"] = [d for d in docs if d in old_docs and old_docs[d] != docs[d]]
    res["docs"]["added"] = [d for d in docs if d not in old_docs]
    res["docs"]["removed"] = [d for d in old_docs if d not in docs]
    res["common_changed"] = part.get("common_hash") != plan["common_hash"]

    old_flows = {k: v for k, v in (part.get("flows") or {}).items() if v.get("status", "active") == "active"}
    if not plan["splittable"] and part.get("plan_hash") != plan["plan_hash"]:
        res["flows"]["changed"] = [f["flow_key"] for f in plan["flows"]] or list(old_flows)
    else:
        for f in plan["flows"]:
            o = old_flows.get(f["flow_key"])
            if o is None:
                res["flows"]["new"].append(f["flow_key"])
            elif o.get("section_hash") != f["section_hash"]:
                res["flows"]["changed"].append(f["flow_key"])
            else:
                res["flows"]["unchanged"].append(f["flow_key"])
        current = {f["flow_key"] for f in plan["flows"]}
        res["flows"]["removed"] = [k for k in old_flows if k not in current]

    docs_changed = any(res["docs"].values())
    structural = docs_changed or res["common_changed"] or res["flows"]["new"] or res["flows"]["changed"] \
        or res["flows"]["removed"]

    drift = None
    if a.drift:
        drift = json.load(sys.stdin) if a.drift == "-" else load_json(a.drift, {})
    if drift is not None:
        for k in res["flows"]["unchanged"]:
            old = {p["screen"]: p for p in old_flows.get(k, {}).get("screens", [])}
            seen, drifted = set(), []
            for entry in drift.get(k, []):
                name = entry.get("screen")
                seen.add(name)
                if entry.get("unreachable"):
                    drifted.append(name)
                    continue
                fp = entry.get("fingerprint") or fingerprint(entry.get("structure", {}))
                if name not in old or old[name].get("fingerprint") != fp:
                    drifted.append(name)
            drifted += [n for n in old if n not in seen]
            if drifted:
                res["screens_drifted"][k] = drifted

    for k in res["flows"]["new"] + res["flows"]["changed"]:
        res["actions"].append({"flow_key": k, "action": "full-explore"})
    for k in res["flows"]["removed"]:
        res["actions"].append({"flow_key": k, "action": "archive"})
    for k in res["flows"]["unchanged"]:
        screens = [p["screen"] for p in old_flows.get(k, {}).get("screens", [])]
        if k in res["screens_drifted"]:
            res["actions"].append({"flow_key": k, "action": "recapture-screens", "screens": res["screens_drifted"][k]})
        elif drift is None and app_changed:
            res["actions"].append({"flow_key": k, "action": "light-check", "screens": screens,
                                   "reached_by": {p["screen"]: p.get("reached_by", "")
                                                  for p in old_flows.get(k, {}).get("screens", [])}})
        else:
            res["actions"].append({"flow_key": k, "action": "none"})

    bits = []
    if docs_changed:
        bits.append("project documents changed")
    if res["common_changed"]:
        bits.append("plan common part changed")
    for n in ("new", "changed", "removed"):
        if res["flows"][n]:
            bits.append(f"flows {n}: {', '.join(res['flows'][n])}")
    if app_changed:
        bits.append(f"app build changed ({part.get('app_version') or '?'} build {part.get('app_build') or '?'} -> "
                    f"{app.get('version') or '?'} build {app.get('build') or '?'})")
    if res["screens_drifted"]:
        bits.append("screens changed in: " + ", ".join(res["screens_drifted"]))

    if structural:
        res["decision"] = "partial"
        res["reason"] = "; ".join(bits)
        res["needs_device"] = any(x["action"] != "archive" and x["action"] != "none" for x in res["actions"])
    elif not app_changed:
        res["decision"] = "stop"
        res["needs_device"] = False
        res["reason"] = "No changes detected, existing test cases are still valid."
    elif drift is None:
        res["decision"] = "check-drift"
        res["reason"] = ("documents and plan unchanged, app build changed: run the light check of the unchanged flows "
                         "and call compare again with --drift")
    elif res["screens_drifted"]:
        res["decision"] = "partial"
        res["reason"] = "; ".join(bits)
    else:
        res["decision"] = "stop"
        res["update_app_only"] = True
        res["reason"] = ("No changes detected, existing test cases are still valid. The new app build was checked and "
                         "no screen changed.")
    dump(res)


# ---------------------------------------------------------------- archive

def _move(src, dest):
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    if os.path.exists(dest):
        root, ext = os.path.splitext(dest)
        n = 2
        while os.path.exists(f"{root}_{n}{ext}"):
            n += 1
        dest = f"{root}_{n}{ext}"
    shutil.move(src, dest)
    return dest.replace("\\", "/")


def archive_flow(flow_dir, ts, files=None, all_files=False):
    moved = []
    hist = os.path.join(flow_dir, "history", ts)
    if all_files:
        targets = [os.path.join(flow_dir, n) for n in os.listdir(flow_dir) if n != "history"]
    else:
        targets = []
        for pat in files or []:
            targets += glob.glob(os.path.join(flow_dir, pat))
    for t in targets:
        if os.path.isdir(t):
            for dp, _, fs in os.walk(t):
                for f in fs:
                    p = os.path.join(dp, f)
                    moved.append(_move(p, os.path.join(hist, os.path.relpath(p, flow_dir))))
            shutil.rmtree(t, ignore_errors=True)  # only empty dirs remain after the moves
            os.makedirs(t, exist_ok=True)
        elif os.path.isfile(t):
            moved.append(_move(t, os.path.join(hist, os.path.relpath(t, flow_dir))))
    return moved


def cmd_archive(a):
    ts = a.ts or ist_file_ts()
    out = {"timestamp": ts, "moved": [], "copied": []}
    if a.plan_dir:
        pd = a.plan_dir
        for plat in (a.platform or PLATFORMS):
            pdir = os.path.join(pd, plat)
            if not os.path.isdir(pdir):
                continue
            for n in sorted(os.listdir(pdir)):
                p = os.path.join(pdir, n)
                if os.path.isdir(p) and n not in SKIP_DIRS:
                    out["moved"] += archive_flow(p, ts, all_files=True)
        pu = os.path.join(pd, "project-understanding.md")
        if os.path.isfile(pu) and not a.keep_understanding:
            out["moved"].append(_move(pu, os.path.join(pd, "history", ts, "project-understanding.md")))
        mf = os.path.join(pd, "manifest.json")
        if os.path.isfile(mf):  # kept in place (holds versions and run history); a copy is archived
            dest = os.path.join(pd, "history", ts, "manifest.json")
            os.makedirs(os.path.dirname(dest), exist_ok=True)
            shutil.copy2(mf, dest)
            out["copied"].append(dest.replace("\\", "/"))
    elif a.flow_dir:
        if not (a.all or a.files):
            sys.exit("archive --flow-dir needs --all or --files")
        out["moved"] = archive_flow(a.flow_dir, ts, a.files, a.all)
    else:
        sys.exit("archive needs --flow-dir or --plan-dir")
    dump(out)


# ---------------------------------------------------------------- scrub

SENSITIVE_HEADERS = {"authorization", "cookie", "set-cookie", "proxy-authorization", "x-api-key",
                     "x-auth-token", "x-csrf-token", "x-xsrf-token", "x-access-token"}
SENSITIVE_KEY = re.compile(r"(pass(word|wd)?$|^pwd$|token|secret|otp|api[-_]?key|authorization|cookie|"
                           r"session|credential|jwt|bearer|^sid$|^auth$)", re.I)


def _scrub_obj(o):
    if isinstance(o, dict):
        return {k: (MASK if SENSITIVE_KEY.search(str(k)) and not isinstance(v, (dict, list)) else _scrub_obj(v))
                for k, v in o.items()}
    if isinstance(o, list):
        return [_scrub_obj(x) for x in o]
    return o


def _scrub_text_body(text):
    if not text:
        return text
    try:
        return json.dumps(_scrub_obj(json.loads(text)), ensure_ascii=False)
    except (ValueError, TypeError):
        pass
    return re.sub(r"(?i)\b([\w.-]*(?:password|passwd|pwd|token|secret|otp|api_?key|session)[\w.-]*)=([^&\s]*)",
                  lambda m: f"{m.group(1)}={MASK}", text)


def _scrub_url(u):
    return re.sub(r"(?i)([?&][\w.-]*(?:password|pwd|token|secret|otp|api_?key|session|code)[\w.-]*=)[^&#]*",
                  lambda m: m.group(1) + MASK, u)


def scrub_har(har):
    for e in har.get("log", {}).get("entries", []):
        for part in ("request", "response"):
            r = e.get(part) or {}
            for h in r.get("headers", []) or []:
                if str(h.get("name", "")).lower() in SENSITIVE_HEADERS or SENSITIVE_KEY.search(str(h.get("name", ""))):
                    h["value"] = MASK
            for c in r.get("cookies", []) or []:
                c["value"] = MASK
            for q in r.get("queryString", []) or []:
                if SENSITIVE_KEY.search(str(q.get("name", ""))) or str(q.get("name", "")).lower() == "code":
                    q["value"] = MASK
            if "url" in r:
                r["url"] = _scrub_url(r["url"])
            pd = r.get("postData")
            if pd:
                if "text" in pd:
                    pd["text"] = _scrub_text_body(pd["text"])
                for p in pd.get("params", []) or []:
                    if SENSITIVE_KEY.search(str(p.get("name", ""))):
                        p["value"] = MASK
            content = r.get("content")
            if content and isinstance(content.get("text"), str) and content.get("encoding") != "base64":
                content["text"] = _scrub_text_body(content["text"])
            if part == "response" and r.get("redirectURL"):
                r["redirectURL"] = _scrub_url(r["redirectURL"])
    return har



def scrub_text(text, secrets):
    kind = "text"
    try:
        data = json.loads(text)
        if isinstance(data, dict) and "log" in data:
            data, kind = scrub_har(data), "har"
        else:
            data, kind = _scrub_obj(data), "json"
        text = json.dumps(data, indent=2, ensure_ascii=False)
    except ValueError:
        text = _scrub_text_body(text) if "\n" not in text.strip() else "\n".join(
            _scrub_text_body(l) for l in text.split("\n"))
    text = re.sub(r"(?i)\b(Bearer|Basic)\s+[A-Za-z0-9\-._~+/=]{8,}", lambda m: f"{m.group(1)} {MASK}", text)
    text = re.sub(r"\beyJ[A-Za-z0-9_-]{5,}\.[A-Za-z0-9_-]{5,}\.[A-Za-z0-9_-]{5,}\b", MASK, text)  # JWTs
    text = re.sub(r"(?i)\b((?:password|passwd|pwd|token|secret|otp|api_?key|session|cookie)[\w.-]*)\s*[:=]\s*(\"?)[^\s\",;&]+\2",
                  lambda m: f"{m.group(1)}={MASK}", text)
    from urllib.parse import quote
    for s in secrets or []:
        if s:
            for v in {s, quote(s, safe=""), json.dumps(s)[1:-1]}:
                text = text.replace(v, MASK)
    return text, kind


def cmd_scrub(a):
    with open(a.file, encoding="utf-8", errors="replace") as f:
        text = f.read()
    text, kind = scrub_text(text, a.secret)
    out = a.out or a.file
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)
    dump({"scrubbed": out, "kind": kind, "secrets_masked": len(a.secret or [])})


# ---------------------------------------------------------------- manifest

def _write_manifest(plan_dir, manifest):
    os.makedirs(plan_dir, exist_ok=True)
    mpath = os.path.join(plan_dir, "manifest.json")
    tmp = mpath + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="\n") as f:
        json.dump(manifest, f, indent=2, ensure_ascii=False)
    os.replace(tmp, mpath)
    return mpath


def cmd_update_manifest(a):
    plan_dir = a.plan_dir
    root = os.path.abspath(a.root or os.path.dirname(os.path.dirname(os.path.abspath(plan_dir))))
    old = load_manifest(plan_dir) or {}
    prev = platform_part(old, a.platform) or {}
    plan = parse_plan(a.test_plan)
    app = app_info(a.app, a.platform, a.serial)
    stories = {}
    for item in a.story or []:  # "<flow-key>=<ref>|<title>" (ref or title may be empty)
        key, _, rest = item.partition("=")
        ref, _, title = rest.partition("|")
        stories[key.strip()] = {"user_story_ref": ref.strip(), "user_story_title": title.strip()}
    flows, current = {}, set()
    for f in plan["flows"]:
        current.add(f["flow_key"])
        fd = os.path.join(plan_dir, a.platform, f["flow_key"])
        screens = flow_screens(fd) if os.path.isdir(fd) else []
        pf = (prev.get("flows") or {}).get(f["flow_key"], {})
        if not screens:  # keep the previous screen list if this run captured nothing new for the flow
            screens = pf.get("screens", [])
        flows[f["flow_key"]] = {"flow_key": f["flow_key"], "name": f["name"], "section_hash": f["section_hash"],
                                "status": "active", "screens": screens}
        story = stories.get(f["flow_key"]) or {k: pf[k] for k in ("user_story_ref", "user_story_title") if pf.get(k)}
        for k, v in story.items():
            if v:  # only what the plan actually names; never invent a reference
                flows[f["flow_key"]][k] = v
    for k, v in (prev.get("flows") or {}).items():
        if k not in current:
            v = dict(v)
            v["status"] = "archived"
            v.setdefault("archived_ist", ist_cell_ts())
            flows[k] = v
    rel = lambda p: os.path.relpath(os.path.abspath(p), root).replace("\\", "/") if p else None
    now = ist_cell_ts()
    runs = list(prev.get("runs", []))
    runs.append({"at_ist": now, "mode": a.mode, "output": rel(a.output_file), "version": a.version,
                 "summary": a.summary or "", "app_version": app.get("version") or prev.get("app_version", ""),
                 "app_build": app.get("build") or prev.get("app_build", ""), "device": a.device or prev.get("device", "")})
    outputs = list(prev.get("outputs", []))
    for p in [a.output_file] + list(a.extra_output or []):
        if p:
            outputs.append({"file": rel(p), "at_ist": now, "version": a.version})
    app_path = app.get("path") if app.get("kind") == "file" else None
    part = {
        "platform": a.platform,
        "app_file_path": rel(app_path) if app_path else (prev.get("app_file_path") if not a.app else None),
        "app_id": app.get("package") or prev.get("app_id", ""),
        "app_hash": app.get("fingerprint") or prev.get("app_hash", ""),
        "app_version": app.get("version") or prev.get("app_version", ""),
        "app_build": app.get("build") or prev.get("app_build", ""),
        "device": a.device or prev.get("device", ""),
        "os_version": a.os_version or prev.get("os_version", ""),
        "flutter_layer": a.flutter_layer if a.flutter_layer is not None else prev.get("flutter_layer"),
        "plan_hash": plan["plan_hash"],
        "common_hash": plan["common_hash"],
        "doc_hashes": list_docs(a.project_dir),
        "flows": flows,
        "last_output_file": rel(a.output_file) or prev.get("last_output_file"),
        "version": a.version,
        "runs": runs,
        "outputs": outputs,
    }
    manifest = {
        "plan_key": os.path.basename(os.path.normpath(plan_dir)),
        "plan_name": plan["plan_name"],
        "plan_file": rel(a.test_plan),
        "test_plan_path": rel(a.test_plan),
        "platforms": dict(old.get("platforms") or {}),
    }
    manifest["platforms"][a.platform] = part
    mpath = _write_manifest(plan_dir, manifest)
    dump({"manifest": mpath, "platform": a.platform, "flows": {k: v["status"] for k, v in flows.items()},
          "version": a.version, "app_hash": part["app_hash"]})


def cmd_update_app(a):
    """A new build whose light check found no changed screen: record the build, keep everything else."""
    manifest = load_manifest(a.plan_dir)
    part = platform_part(manifest, a.platform)
    if not part:
        sys.exit(f"no {a.platform} part in manifest.json")
    app = app_info(a.app, a.platform, a.serial)
    if not app.get("fingerprint"):
        sys.exit("the app build fingerprint could not be computed")
    old = (part.get("app_version"), part.get("app_build"))
    part.update(app_hash=app["fingerprint"], app_version=app.get("version") or part.get("app_version", ""),
                app_build=app.get("build") or part.get("app_build", ""))
    if a.device:
        part["device"] = a.device
    for fk, screens in (load_json(a.drift, {}) if a.drift and a.drift != "-" else
                        (json.load(sys.stdin) if a.drift == "-" else {})).items():
        flow = (part.get("flows") or {}).get(fk)
        if not flow:
            continue
        for s in flow.get("screens", []):
            if any(e.get("screen") == s["screen"] and not e.get("unreachable") for e in screens):
                s["last_checked_ist"] = ist_cell_ts()
    part.setdefault("runs", []).append({"at_ist": ist_cell_ts(), "mode": "light-check", "output": None,
                                        "version": part.get("version"), "app_version": part["app_version"],
                                        "app_build": part["app_build"],
                                        "summary": f"New app build {old[0]} ({old[1]}) -> {part['app_version']} "
                                                   f"({part['app_build']}); light check found no changed screen"})
    _write_manifest(a.plan_dir, manifest)
    dump({"ok": True, "platform": a.platform, "app_hash": part["app_hash"], "app_version": part["app_version"],
          "app_build": part["app_build"]})


def cmd_fingerprint(a):
    out = {}
    for p in a.files:
        s = load_json(p, {})
        fp = fingerprint(s)
        out[p] = fp
        if a.write and s.get("fingerprint") != fp:
            s["fingerprint"] = fp
            with open(p, "w", encoding="utf-8", newline="\n") as f:
                json.dump(s, f, indent=2, ensure_ascii=False)
    dump(out)


# ---------------------------------------------------------------- main

def main():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)

    p = sp.add_parser("plan", help="split the test plan into flows")
    p.add_argument("--test-plan", required=True)
    p.set_defaults(fn=cmd_plan)

    def app_args(p, required_platform=True):
        p.add_argument("--platform", choices=PLATFORMS, required=required_platform)
        p.add_argument("--app", help="APK/IPA path or package/bundle id from the test plan")
        p.add_argument("--serial", help="real device serial/UDID (to fingerprint an installed app)")

    p = sp.add_parser("compare", help="decide full / partial / check-drift / stop for one platform")
    p.add_argument("--project-dir", required=True)
    p.add_argument("--test-plan", required=True)
    p.add_argument("--plan-dir", required=True, help="Exploration/<plan-key>")
    app_args(p)
    p.add_argument("--drift", help="light-check JSON file, or - for stdin")
    p.add_argument("--force-full", action="store_true")
    p.set_defaults(fn=cmd_compare)

    p = sp.add_parser("fingerprint", help="fingerprint screen JSON files")
    p.add_argument("files", nargs="+")
    p.add_argument("--write", action="store_true", help="store the fingerprint inside each file")
    p.set_defaults(fn=cmd_fingerprint)

    p = sp.add_parser("archive", help="move files into history/<IST-timestamp>/")
    p.add_argument("--flow-dir")
    p.add_argument("--plan-dir", help="archive every flow (of --platform, or all) + project-understanding (--force-full)")
    p.add_argument("--platform", action="append", choices=PLATFORMS)
    p.add_argument("--keep-understanding", action="store_true")
    p.add_argument("--files", nargs="+", help="paths/globs relative to --flow-dir")
    p.add_argument("--all", action="store_true")
    p.add_argument("--ts", help="timestamp folder name (default: now, IST)")
    p.set_defaults(fn=cmd_archive)

    for name in ("scrub", "scrub-log", "scrub-har"):
        p = sp.add_parser(name, help="mask secrets in a device log / JSON / text file")
        p.add_argument("file")
        p.add_argument("--secret", action="append", help="literal value to mask everywhere (repeatable)")
        p.add_argument("--out", help="write here instead of in place")
        p.set_defaults(fn=cmd_scrub)

    p = sp.add_parser("update-manifest", help="write the platform part of manifest.json after a run")
    p.add_argument("--project-dir", required=True)
    p.add_argument("--test-plan", required=True)
    p.add_argument("--plan-dir", required=True)
    app_args(p)
    p.add_argument("--device", help="model of the real device, e.g. 'Pixel 7'")
    p.add_argument("--os-version", help="e.g. 'Android 14'")
    p.add_argument("--flutter-layer", type=lambda v: v.lower() in ("1", "true", "yes"), default=None)
    p.add_argument("--output-file", required=True)
    p.add_argument("--version", required=True)
    p.add_argument("--mode", required=True, choices=["first", "partial", "force-full"])
    p.add_argument("--summary", default="")
    p.add_argument("--extra-output", action="append", help="another generated file of this run (e.g. the .xlsx)")
    p.add_argument("--story", action="append", help='user story of a flow: "<flow-key>=<ref>|<title>"; repeatable')
    p.add_argument("--root", help="project root (default: parent of Exploration/)")
    p.set_defaults(fn=cmd_update_manifest)

    p = sp.add_parser("update-app", help="record a new app build after a light check without changes")
    p.add_argument("--plan-dir", required=True)
    app_args(p)
    p.add_argument("--device")
    p.add_argument("--drift", help="the light-check JSON (file or -), to stamp last_checked_ist")
    p.set_defaults(fn=cmd_update_app)

    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
